"""Intelligent CSV/Excel import and export for list views."""
from __future__ import annotations

import csv
import io
import json
import re
from datetime import date, datetime
from typing import Any

from pyvelm.fields import (
    Boolean,
    Char,
    Date,
    Datetime,
    Float,
    Integer,
    Many2one,
    Monetary,
    One2many,
    Many2many,
    Text,
)

_IMPORTABLE = (
    Boolean,
    Char,
    Text,
    Integer,
    Float,
    Monetary,
    Date,
    Datetime,
    Many2one,
)

_MAX_IMPORT_ROWS = 2000
_PREVIEW_ROWS = 5


def _normalize_key(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", (value or "").lower())


def _humanize_field_label(fname: str) -> str:
    """Turn ``country_id`` into ``Country`` when no field string is set."""
    base = fname
    if base.endswith("_id"):
        base = base[:-3]
    return base.replace("_", " ").strip().title() or fname


def _m2o_lookup_field_names(env, comodel: str) -> list[str]:
    """Search keys tried when resolving a Many2one cell value."""
    if comodel not in env.registry:
        return ["name"]
    cls = env.registry[comodel]
    keys: list[str] = []
    for candidate in ("name", "code", "login", "email"):
        if candidate in cls._fields:
            keys.append(candidate)
    return keys or ["name"]


def m2o_import_hint(env, comodel: str | None) -> str:
    """Short guidance for spreadsheet cells (shown in the import wizard)."""
    if not comodel:
        return "Related record ID or name"
    keys = _m2o_lookup_field_names(env, comodel)
    parts = ["numeric ID"]
    if "name" in keys:
        parts.append("exact name")
    if "code" in keys:
        parts.append("code (e.g. country ISO)")
    return " or ".join(parts)


def list_importable_fields(env, model: str) -> list[dict[str, Any]]:
    """Stored, writable scalar fields suitable for tabular import."""
    if model not in env.registry:
        return []
    cls = env.registry[model]
    out: list[dict[str, Any]] = []
    for fname, field in sorted(cls._fields.items()):
        if fname == "id":
            out.append({
                "name": "id",
                "label": "ID",
                "type": "Integer",
                "required": False,
                "import_hint": "Existing record ID (for updates)",
            })
            continue
        if not isinstance(field, _IMPORTABLE):
            continue
        if field.compute or field.readonly or not field.is_stored:
            continue
        comodel = getattr(field, "comodel_name", None)
        if isinstance(field, Many2one):
            try:
                env.check_access(comodel, "read")
            except PermissionError:
                continue
        label = field.string or _humanize_field_label(fname)
        spec: dict[str, Any] = {
            "name": fname,
            "label": label,
            "type": type(field).__name__,
            "required": bool(field.required),
            "comodel": comodel,
            "choices": getattr(field, "choices", None),
        }
        if isinstance(field, Many2one):
            spec["import_hint"] = m2o_import_hint(env, comodel)
        out.append(spec)
    return out


def filter_import_fields(
    fields: list[dict[str, Any]],
    selected_names: list[str] | None,
) -> list[dict[str, Any]]:
    """Keep only user-selected template columns (default: all)."""
    if not selected_names:
        return list(fields)
    allowed = {n.strip() for n in selected_names if n and str(n).strip()}
    if not allowed:
        return list(fields)
    lookup = {spec["name"]: spec for spec in fields}
    ordered: list[dict[str, Any]] = []
    for name in selected_names:
        key = name.strip()
        if key in lookup and key not in {s["name"] for s in ordered}:
            ordered.append(lookup[key])
    for spec in fields:
        if spec["name"] in allowed and spec["name"] not in {s["name"] for s in ordered}:
            ordered.append(spec)
    return ordered


def parse_fields_query(raw: str | list[str] | None) -> list[str]:
    """Parse ``fields`` from repeated query params or a comma-separated string."""
    if not raw:
        return []
    if isinstance(raw, str):
        parts = raw.split(",")
    else:
        parts = list(raw)
    return [p.strip() for p in parts if p and str(p).strip()]


def suggest_column_mapping(
    headers: list[str],
    fields: list[dict[str, Any]],
) -> dict[int, str]:
    """Map file column indexes to model field names by header heuristics."""
    lookup: dict[str, str] = {}
    for spec in fields:
        lookup[_normalize_key(spec["name"])] = spec["name"]
        lookup[_normalize_key(spec["label"])] = spec["name"]
    mapping: dict[int, str] = {}
    for idx, header in enumerate(headers):
        key = _normalize_key(header)
        if key in lookup:
            mapping[idx] = lookup[key]
    return mapping


def parse_tabular_upload(content: bytes, filename: str) -> tuple[list[str], list[list[Any]]]:
    """Parse CSV or Excel upload into headers + data rows."""
    lower = (filename or "").lower()
    if lower.endswith((".xlsx", ".xlsm", ".xltx", ".xltm")):
        return _parse_xlsx(content)
    return _parse_csv(content)


def _parse_csv(content: bytes) -> tuple[list[str], list[list[Any]]]:
    text = content.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    if not rows:
        return [], []
    headers = [str(h).strip() for h in rows[0]]
    data = [[_clean_cell(c) for c in row] for row in rows[1:] if any(str(c).strip() for c in row)]
    return headers, data


def _parse_xlsx(content: bytes) -> tuple[list[str], list[list[Any]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError(
            "openpyxl is required for Excel import (pip install openpyxl)"
        ) from exc
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    raw_rows = list(ws.iter_rows(values_only=True))
    if not raw_rows:
        return [], []
    headers = [str(h or "").strip() for h in raw_rows[0]]
    data = []
    for row in raw_rows[1:]:
        cells = [_clean_cell(c) for c in row]
        if any(str(c).strip() for c in cells):
            data.append(cells)
    return headers, data


def _clean_cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, bool):
        return value
    if isinstance(value, (datetime, date)):
        return value.isoformat(sep=" ", timespec="seconds")
    return value


def _record_by_id(Model, record_id: int):
    """Singleton recordset for *record_id*, or empty if not in the database."""
    return Model.search([("id", "=", int(record_id))], limit=1)


def _resolve_m2o(env, comodel: str, raw: Any) -> int | None:
    if raw is None or raw == "" or raw is False:
        return None
    Model = env[comodel]
    text = str(raw).strip()
    if text.isdigit():
        found = _record_by_id(Model, int(text))
        if found:
            return found.id
    for key in _m2o_lookup_field_names(env, comodel):
        field = Model._fields.get(key)
        if field is None:
            continue
        # Exact match first (codes, names, logins).
        found = Model.search([(key, "=", text)], limit=1)
        if found:
            return found.id
        if key == "code" and text.upper() != text:
            found = Model.search([(key, "=", text.upper())], limit=1)
            if found:
                return found.id
        if key == "name":
            found = Model.search([(key, "ilike", text)], limit=1)
            if found:
                return found.id
    raise ValueError(
        f"Cannot resolve {comodel} reference {text!r} — "
        f"use a record ID or a matching {', '.join(_m2o_lookup_field_names(env, comodel))}"
    )


def _coerce_field_value(env, model: str, fname: str, raw: Any) -> Any:
    cls = env.registry[model]
    field = cls._fields.get(fname)
    if field is None:
        raise ValueError(f"Unknown field {fname!r}")
    if isinstance(field, Many2one):
        return _resolve_m2o(env, field.comodel_name, raw)
    if isinstance(field, Boolean):
        if raw is None or raw == "":
            return False
        if isinstance(raw, bool):
            return raw
        text = str(raw).strip().lower()
        if text in ("1", "true", "yes", "y", "on"):
            return True
        if text in ("0", "false", "no", "n", "off"):
            return False
        raise ValueError(f"Invalid boolean value {raw!r}")
    if isinstance(field, Char) and getattr(field, "choices", None):
        if raw is None or raw == "":
            return None
        text = str(raw).strip()
        for value, label in field.choices or []:
            if text == value or text == label:
                return value
        raise ValueError(f"Invalid selection {text!r} for {fname}")
    return field.to_python(raw)


def build_row_vals(
    env,
    model: str,
    row: list[Any],
    mapping: dict[int, str],
) -> dict[str, Any]:
    vals: dict[str, Any] = {}
    for col_idx, fname in mapping.items():
        if not fname or fname == "_skip":
            continue
        if col_idx >= len(row):
            continue
        raw = row[col_idx]
        if raw is None or raw == "":
            continue
        vals[fname] = _coerce_field_value(env, model, fname, raw)
    return vals


def import_rows(
    env,
    model: str,
    rows: list[list[Any]],
    mapping: dict[int, str],
    *,
    update_by_id: bool = False,
    atomic: bool = True,
) -> dict[str, Any]:
    """Create or update records from parsed tabular rows.

    When *atomic* is true (default), any row error aborts the whole batch
    by raising :class:`ImportBatchError` so the caller's transaction rolls
    back and no rows are persisted.
    """
    env.check_access(model, "create")
    Model = env[model]
    created = 0
    updated = 0
    errors: list[dict[str, Any]] = []
    for line_no, row in enumerate(rows, start=2):
        try:
            vals = build_row_vals(env, model, row, mapping)
            if not vals:
                continue
            record_id = vals.pop("id", None)
            if update_by_id and record_id:
                env.check_access(model, "write")
                rec = _record_by_id(Model, int(record_id))
                if not rec:
                    raise ValueError(f"Record id={record_id} not found")
                rec.write(vals)
                updated += 1
            else:
                Model.create(vals)
                created += 1
        except Exception as exc:  # noqa: BLE001
            errors.append({"line": line_no, "error": str(exc)})
    result = {
        "created": created,
        "updated": updated,
        "errors": errors,
        "total": len(rows),
    }
    if atomic and errors:
        raise ImportBatchError(result)
    return result


def import_data_line_no(row_index: int) -> int:
    """Spreadsheet line number for a zero-based data row (line 1 = headers)."""
    return row_index + 2


def import_errors_by_line(errors: list[dict[str, Any]] | None) -> dict[int, str]:
    """Map spreadsheet line numbers to messages from :func:`import_rows`."""
    out: dict[int, str] = {}
    for entry in errors or []:
        line = entry.get("line")
        if line is None:
            continue
        out[int(line)] = str(entry.get("error") or "")
    return out


def failed_import_sheet(
    headers: list[str],
    rows: list[list[Any]],
    errors: list[dict[str, Any]],
) -> tuple[list[str], list[list[Any]]]:
    """Build tabular failed rows with an appended Error column."""
    errors_by_line = import_errors_by_line(errors)
    if not errors_by_line:
        return list(headers) + ["Error"], []
    out_headers = list(headers) + ["Error"]
    out_rows: list[list[Any]] = []
    col_count = len(headers)
    for idx, row in enumerate(rows):
        line_no = import_data_line_no(idx)
        err = errors_by_line.get(line_no)
        if not err:
            continue
        cells = list(row)
        if len(cells) < col_count:
            cells.extend([""] * (col_count - len(cells)))
        elif len(cells) > col_count:
            cells = cells[:col_count]
        out_rows.append(cells + [err])
    return out_headers, out_rows


def failed_import_xlsx_bytes(
    headers: list[str],
    rows: list[list[Any]],
    errors: list[dict[str, Any]],
    *,
    title: str = "Failed rows",
) -> bytes:
    """Excel workbook of rows that failed import, with an Error column."""
    sheet_headers, sheet_rows = failed_import_sheet(headers, rows, errors)
    return import_template_xlsx_bytes(
        sheet_headers, title=title, rows=sheet_rows,
    )


def export_list_data(
    env,
    model: str,
    fields_spec: list[dict],
    records,
) -> tuple[list[str], list[list[Any]]]:
    """Build export headers and rows for a recordset."""
    cls = env.registry[model]
    columns: list[tuple[str, str]] = []
    for spec in fields_spec:
        fname = spec["name"]
        if fname not in cls._fields:
            continue
        field = cls._fields[fname]
        if isinstance(field, (One2many, Many2many)):
            continue
        columns.append((fname, spec.get("label") or field.string or fname))

    headers = [label for _fname, label in columns]
    rows: list[list[Any]] = []
    for rec in records:
        row: list[Any] = []
        for fname, _label in columns:
            row.append(_export_cell(getattr(rec, fname, None)))
        rows.append(row)
    return headers, rows


def _export_cell(value: Any) -> Any:
    if value is None or value is False:
        return ""
    if hasattr(value, "_ids"):
        if len(value._ids) == 1:
            rec = value
            if hasattr(rec, "display_name"):
                return rec.display_name
            if hasattr(rec, "name"):
                return rec.name
            return value._ids[0]
        return ",".join(str(i) for i in value._ids)
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, default=str)
    return value


def import_template_fields_for_view(
    env,
    view,
    selected_names: list[str] | None = None,
) -> list[dict[str, Any]]:
    """Importable fields for a template, optionally filtered by user selection."""
    importable = list_importable_fields(env, view.model)
    return filter_import_fields(importable, selected_names)


def import_template_headers(fields: list[dict[str, Any]]) -> list[str]:
    """Column headers for a blank import template (Odoo-style labels)."""
    return [spec.get("label") or spec["name"] for spec in fields]


def import_template_data_rows(
    env,
    model: str,
    fields: list[dict[str, Any]],
    records,
) -> list[list[Any]]:
    """Build template data rows for *records* using the selected import columns."""
    fields_spec = [
        {"name": f["name"], "label": f.get("label") or f["name"]}
        for f in fields
    ]
    _headers, rows = export_list_data(env, model, fields_spec, records)
    return rows


def import_template_xlsx_bytes(
    headers: list[str],
    *,
    title: str = "Import",
    rows: list[list[Any]] | None = None,
) -> bytes:
    """Excel import template via openpyxl (headers only, or with optional data rows)."""
    return export_xlsx_bytes(headers, rows or [], title=title)


def parse_include_data_query(value: str | None) -> bool:
    """True when the user asked to pre-fill the template with existing records."""
    return str(value or "").strip().lower() in ("1", "true", "yes", "on")


def export_csv_bytes(headers: list[str], rows: list[list[Any]]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(headers)
    for row in rows:
        writer.writerow(row)
    return buf.getvalue().encode("utf-8-sig")


def export_xlsx_bytes(headers: list[str], rows: list[list[Any]], *, title: str = "Export") -> bytes:
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise RuntimeError(
            "openpyxl is required for Excel export (pip install openpyxl)"
        ) from exc
    wb = Workbook()
    ws = wb.active
    ws.title = (title[:31] if title else "Export") or "Export"
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def encode_import_payload(headers: list[str], rows: list[list[Any]]) -> str:
    """JSON payload for the preview → run step (capped)."""
    capped = rows[:_MAX_IMPORT_ROWS]
    return json.dumps({"headers": headers, "rows": capped}, default=str)


def decode_import_payload(raw: str) -> tuple[list[str], list[list[Any]]]:
    data = json.loads(raw)
    return list(data.get("headers") or []), list(data.get("rows") or [])


def mapping_from_form(
    form,
    headers: list[str],
    fields: list[dict[str, Any]],
) -> dict[int, str]:
    """Read column mapping selects from the import preview form."""
    mapping: dict[int, str] = {}
    for key, value in form.items():
        if not str(key).startswith("map_"):
            continue
        try:
            idx = int(str(key)[4:])
        except ValueError:
            continue
        if value:
            mapping[idx] = str(value)
    if not mapping:
        mapping = suggest_column_mapping(headers, fields)
    return mapping


class ImportBatchError(Exception):
    """Raised when an atomic import has row errors — rolls back the transaction."""

    def __init__(self, result: dict[str, Any]):
        self.result = result
        super().__init__(
            f"{len(result.get('errors') or [])} import error(s) — batch aborted"
        )


class ImportTestAbort(Exception):
    """Raised to roll back a dry-run import test transaction."""
